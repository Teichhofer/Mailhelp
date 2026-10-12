"""xlsx list of an activity report, built in memory for one download.

Mail content is untrusted: every text cell is stored as a string (never as a
formula), values that a spreadsheet could read as a formula get a leading
apostrophe, control characters are removed and no value exceeds Excel's cell
limit.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .. import __version__
from .models import ActivityReport, ReportEvent, ReportMail

CELL_LIMIT = 32767
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
REMOVED = "(Volltext nach Ablauf der Aufbewahrungsfrist entfernt)"


def safe_text(value: str) -> str:
    text = _CONTROL.sub("", value).replace("\r\n", "\n").replace("\r", "\n")
    if text.startswith(_FORMULA_START):
        text = "'" + text
    return text[:CELL_LIMIT]


def _row(sheet: Worksheet, values: list[Any]) -> None:
    sheet.append([safe_text(value) if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # openpyxl would otherwise store "=…" as a formula


def _header(sheet: Worksheet, titles: list[str], widths: list[int]) -> None:
    _row(sheet, titles)
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"


def _local(moment: datetime, zone: str) -> datetime:
    # Excel has no time zones: the local wall-clock time is stored.
    return moment.astimezone(ZoneInfo(zone)).replace(tzinfo=None)


def _clock(event: ReportEvent) -> str:
    if event.start_time and event.end_time:
        return f"{event.start_time}–{event.end_time}"
    return event.start_time or ""


def _identity(event: ReportEvent) -> tuple[date, str, str]:
    return event.day, event.start_time or "", " ".join(event.title.split()).casefold()


def event_rows(mails: list[ReportMail]) -> list[tuple[int, ReportEvent, ReportMail]]:
    """Exported dates in date order; the same date from several mails shares one number."""
    rows = [(event, mail) for mail in mails if mail.counts_as_relevant and mail.skipped is None
            for event in mail.events if not event.excluded and not event.cancelled]
    rows.sort(key=lambda item: (item[0].day, item[0].start_time or "", item[1].received_at))
    numbers: dict[tuple[date, str, str], int] = {}
    return [(numbers.setdefault(_identity(event), len(numbers) + 1), event, mail) for event, mail in rows]


def without_events(mails: list[ReportMail]) -> list[ReportMail]:
    return [mail for mail in mails if mail.counts_as_relevant and mail.skipped is None
            and not any(not event.excluded and not event.cancelled for event in mail.events)]


def workbook(report: ActivityReport, mails: list[ReportMail], zone: str, with_unscheduled: bool) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Termine"
    _header(sheet, ["Termin-Nr.", "Termin", "Datum", "Uhrzeit", "Ort", "Mail-Betreff", "Absender",
                    "Mail-Datum", "Volltext"], [11, 40, 12, 13, 25, 40, 35, 17, 80])
    for number, event, mail in event_rows(mails):
        _row(sheet, [number, event.title, event.day, _clock(event), event.location or "", mail.subject,
                     mail.sender, _local(mail.received_at, zone), REMOVED if mail.text_removed else mail.text])
        sheet.cell(sheet.max_row, 3).number_format = "DD.MM.YYYY"
        sheet.cell(sheet.max_row, 8).number_format = "DD.MM.YYYY HH:MM"
    sheet.auto_filter.ref = sheet.dimensions
    if with_unscheduled:
        other = book.create_sheet("Ohne Termin")
        _header(other, ["Mail-Betreff", "Absender", "Mail-Datum", "Volltext"], [40, 35, 17, 80])
        for mail in without_events(mails):
            _row(other, [mail.subject, mail.sender, _local(mail.received_at, zone),
                         REMOVED if mail.text_removed else mail.text])
            other.cell(other.max_row, 3).number_format = "DD.MM.YYYY HH:MM"
    _summary(book.create_sheet("Bericht"), report, mails)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _summary(sheet: Worksheet, report: ActivityReport, mails: list[ReportMail]) -> None:
    definition = report.definition
    corrected = any(mail.user_relevant is not None or any(e.edited or e.manual or e.excluded for e in mail.events)
                    for mail in mails)
    _header(sheet, ["Angabe", "Wert"], [24, 90])
    for label, value in (
            ("Titel", definition.title), ("Thema", definition.topic),
            ("Zeitraum", f"{definition.since:%d.%m.%Y} – {definition.until:%d.%m.%Y}"),
            ("Ordner", ", ".join(item.folder for item in report.folders) or ", ".join(definition.folders)),
            ("Erstellt", f"{report.created_at:%d.%m.%Y %H:%M} UTC"),
            ("Mails gelesen", str(report.counters.read)), ("Relevante Mails", str(report.counters.relevant)),
            ("Termine", str(len(event_rows(mails)))), ("Mailhelp-Version", __version__),
            ("Hinweis", "KI-gestützt erstellt; vom Nutzer geprüft: " + ("ja" if corrected else "nein"))):
        _row(sheet, [label, value])


def filename(report: ActivityReport, today: date) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", report.definition.title).strip("-")[:60] or "Bericht"
    return f"Taetigkeitsbericht-{slug}-{today:%Y-%m-%d}.xlsx"


__all__ = ["CELL_LIMIT", "REMOVED", "event_rows", "filename", "safe_text", "without_events", "workbook"]
