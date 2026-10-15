"""Web pages of activity reports (2.2): authorization, CSRF, versions, idempotency and the xlsx download."""
from __future__ import annotations

import io
import re
from datetime import UTC, date, datetime
from html import unescape
from types import SimpleNamespace as NS

import pytest
from activity_fakes import definition
from openpyxl import load_workbook
from shared_web_auth import Browser

from mailhelp.activity.models import FolderProgress, ReportEvent, ReportMail, ReportStatus
from mailhelp.activity.service import ActivityService
from mailhelp.activity.store import ActivityReportStore
from mailhelp.domain.relevance import IrrelevantSenders

RECEIVED = datetime(2026, 1, 10, 8, 30, tzinfo=UTC)
NEW = {"title": "Elternbeirat", "topic": "Termine des Elternbeirats der Klasse 4b", "since": "2026-01-01",
       "until": "2026-01-31", "folders": "INBOX\r\nArchiv\r\n", "include_sent": "on"}


@pytest.fixture
def reports(web):
    service = ActivityService(ActivityReportStore(web.store), clock=lambda: RECEIVED)
    web.site.activity.service = service
    web.site.activity.suggest_folders = lambda mailbox: ["INBOX", "Schule"]
    web.site.activity.available_folders = lambda mailbox: ["INBOX", "Schule", "Archiv"]
    return service


def page(web, path):
    response = web.browser.request("GET", path)
    return response.status, unescape(response.body.decode())


def post(web, path, fields):
    return web.browser.request("POST", path, fields)


def csrf(web):
    return web.auth.csrf_token(web.auth.sessions()[0].id)


def new_form(web):
    html = web.browser.request("GET", "/berichte/neu").body.decode()
    return dict(re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)">', html))


def mail(uid, subject, events=(), **changes):
    values = dict(key=f"INBOX:7:{uid}", folder="INBOX", uid=uid, received_at=RECEIVED,
                  sender="Schule <info@schule.test>", subject=subject, text=f"Text {subject}", relevant=True,
                  confidence="high", reason="Gehört zum Elternbeirat", events=list(events), events_extracted=True)
    return ReportMail(**(values | changes))


def filled(reports, status=ReportStatus.DONE):
    report = reports.create(definition(), [], [])
    report.status = status
    reports.store.save(report)
    for item in (
            mail(1, "Einladung", [ReportEvent(id="e1", title="Elternabend", day=date(2026, 1, 20),
                                              start_time="19:00", end_time="20:30", location="Aula",
                                              evidence="am 20.01. um 19 Uhr", out_of_range=True),
                                  ReportEvent(id="e2", title="Basar", day=date(2026, 1, 22), manual=True),
                                  ReportEvent(id="e3", title="Sitzung", day=date(2026, 1, 23), edited=True),
                                  ReportEvent(id="e4", title="Alt", day=date(2026, 1, 2), excluded=True)],
                 confidence="low"),
            mail(2, "Protokoll", attachments=["protokoll.pdf"]),
            mail(3, "Werbung", relevant=False, reason="Werbung", events_extracted=False),
            mail(4, "Gesperrt", relevant=None, skipped="sender", events_extracted=False),
            mail(5, "", relevant=None, events_extracted=False),
            mail(6, "Gelöscht", [ReportEvent(id="e1", title="Fest", day=date(2026, 1, 25))],
                 relevant=False, user_relevant=True, text="", text_removed=True)):
        reports.store.save_mail(report.id, item)
    return reports.store.load(report.id)


def test_pages_require_a_session_and_posts_a_csrf_token(web, reports):
    report = filled(reports)
    stranger = Browser(web.app)
    for path in ("/berichte", "/berichte/neu", f"/berichte/{report.id}", f"/berichte/{report.id}/mail/0",
                 f"/berichte/{report.id}/export.xlsx"):
        assert stranger.request("GET", path).headers["Location"] == "/anmelden"
    for path in ("/berichte", f"/berichte/{report.id}/aktion", f"/berichte/{report.id}/mail/0"):
        assert stranger.request("POST", path, {"csrf": "x"}).status == 403
        assert post(web, path, {"csrf": "falsch", "action": "start", "version": "0", "key": "k" * 16}).status == 403
    assert reports.store.load(report.id).version == report.version


def test_switched_off_reports_explain_how_to_enable_them(web):
    assert web.site.activity.service is None
    status, html = page(web, "/berichte")
    assert status == 200 and "activity_report.enabled" in html


    report_id = "a" * 32
    for path in ("/berichte/neu", f"/berichte/{report_id}", f"/berichte/{report_id}/mail/0",
                 f"/berichte/{report_id}/export.xlsx"):
        assert page(web, path)[0] == 404
    for path in ("/berichte", f"/berichte/{report_id}/aktion", f"/berichte/{report_id}/mail/0"):
        response = post(web, path, {"csrf": csrf(web)})
        assert response.status == 404 and b"activity_report.enabled" in response.body


def test_delete_requires_an_idempotency_key(web, reports):
    report = filled(reports)
    response = post(web, f"/berichte/{report.id}/aktion",
                    {"csrf": csrf(web), "version": str(report.version), "action": "delete"})
    assert response.status == 409
    assert reports.store.load(report.id) is not None


def test_report_filters_fingerprint_repeat_and_unknown_folders(web, reports):
    report = filled(reports)
    url = f"/berichte/{report.id}"
    status, html = page(web, url + "?month=2026-01&excluded=on")
    assert status == 200 and "ausgeschlossen" in html
    first = web.browser.request("GET", f"/api/stand?bericht={report.id}").body
    report.message = "Fortschritt"
    reports.store.save(report)
    assert web.browser.request("GET", f"/api/stand?bericht={report.id}").body != first
    assert web.browser.request("GET", "/api/stand?bericht=" + "f" * 32).status == 404
    fields = {"csrf": csrf(web), "key": "repeat-web-001", "version": str(report.version),
              "action": "repeat", "topic": "Ein anderes Thema fuer diese Mails"}
    result = post(web, url + "/aktion", fields)
    assert result.status == 303 and result.headers["Location"] != url
    assert post(web, url + "/aktion", fields).headers == result.headers
    assert post(web, "/berichte", new_form(web) | NEW | {"folders": "Unknown"}).status == 422
    web.site.activity.available_folders = lambda mailbox: (_ for _ in ()).throw(OSError("offline"))
    assert post(web, "/berichte", new_form(web) | NEW).status == 503
    web.site.activity.service = None
    assert web.browser.request("GET", f"/api/stand?bericht={report.id}").status == 404


def test_create_counts_first_and_a_replayed_form_creates_nothing_twice(web, reports):
    status, html = page(web, "/berichte")
    assert status == 200 and "Noch keine Berichte." in html and 'href="/berichte"' in html
    status, html = page(web, "/berichte/neu")
    assert status == 200 and "INBOX\nSchule</textarea>" in html and 'name="include_sent" checked' in html
    web.site.archive.configuration = NS(blocked_senders=lambda: IrrelevantSenders(
        addresses=["news@shop.test"], domains=["werbung.test"]))
    fields = new_form(web) | NEW
    response = post(web, "/berichte", fields)
    assert response.status == 303
    report_id = response.headers["Location"].rsplit("/", 1)[1]
    report = reports.store.load(report_id)
    assert report.status == ReportStatus.COUNTING and report.definition.folders == ["INBOX", "Archiv"]
    assert report.definition.include_sent and not report.definition.skip_automated
    assert (report.blocked_addresses, report.blocked_domains) == (["news@shop.test"], ["werbung.test"])
    # A double click on the button sends the same key again.
    assert post(web, "/berichte", fields).headers["Location"] == f"/berichte/{report_id}"
    assert len(reports.store.reports()) == 1
    assert post(web, "/berichte", fields | {"title": "Anders"}).status == 409
    assert post(web, "/berichte", fields | {"key": "kurz"}).status == 409
    web.site.archive.configuration = None
    other = post(web, "/berichte", new_form(web) | NEW | {"skip_automated": "on"})
    assert reports.store.load(other.headers["Location"].rsplit("/", 1)[1]).blocked_addresses == []
    status, html = page(web, "/berichte")
    assert html.count('<a href="/berichte/') == 2 and "0 von 0 Mails" in html


@pytest.mark.parametrize("change", [{"topic": "kurz"}, {"since": "gestern"}, {"until": "2025-01-01"},
                                    {"folders": "INBOX\nINBOX"}, {"folders": " "}])
def test_invalid_definitions_keep_the_input(web, reports, change):
    response = post(web, "/berichte", new_form(web) | NEW | {"title": "<b>Mein Bericht</b>"} | change)
    html = response.body.decode()
    assert response.status == 422 and "&lt;b&gt;Mein Bericht&lt;/b&gt;" in html and "notice-error" in html
    assert reports.store.reports() == []


def test_unreachable_mailbox_still_offers_the_form(web, reports):
    def broken(_mailbox):
        raise OSError("imap down")
    web.site.activity.suggest_folders = broken
    status, html = page(web, "/berichte/neu")
    assert status == 200 and "INBOX</textarea>" in html
    assert web.logger.events[-1][0][2] == "folder_suggestion_failed"


def test_start_pause_cancel_and_delete_are_version_checked(web, reports):
    report = reports.create(definition(), [], [])
    status, html = page(web, f"/berichte/{report.id}")
    assert status == 200 and "Zählt die Mails" in html and 'http-equiv="refresh"' in html
    assert "kostenpflichtig" not in html
    report.status = ReportStatus.READY
    report.folders = [FolderProgress(folder="INBOX", uids=[1, 2, 3])]
    reports.store.save(report)
    status, html = page(web, f"/berichte/{report.id}")
    assert "3 Mails im Zeitraum" in html and 'http-equiv="refresh"' not in html
    start = _action(web, report.id, "start")
    assert post(web, f"/berichte/{report.id}/aktion", start).headers["Location"] == f"/berichte/{report.id}"
    assert reports.store.load(report.id).status == ReportStatus.RUNNING
    # The same submission again changes nothing; another stale form is refused.
    assert post(web, f"/berichte/{report.id}/aktion", start).status == 303
    assert reports.store.load(report.id).version == report.version + 1
    assert post(web, f"/berichte/{report.id}/aktion", start | {"key": "neu-12345678"}).status == 409
    assert post(web, f"/berichte/{report.id}/aktion", start | {"version": "x"}).status == 422
    delete = _action(web, report.id, "pause") | {"action": "delete"}
    response = post(web, f"/berichte/{report.id}/aktion", delete)
    assert response.status == 409 and "pausiert oder abgebrochen" in response.body.decode()
    post(web, f"/berichte/{report.id}/aktion", _action(web, report.id, "pause"))
    post(web, f"/berichte/{report.id}/aktion", _action(web, report.id, "resume"))
    post(web, f"/berichte/{report.id}/aktion", _action(web, report.id, "cancel"))
    assert reports.store.load(report.id).status == ReportStatus.CANCELLED
    response = post(web, f"/berichte/{report.id}/aktion", _action(web, report.id, "delete"))
    assert response.status == 200 and "gelöscht" in response.body.decode()
    assert reports.store.load(report.id) is None and page(web, f"/berichte/{report.id}")[0] == 404
    assert web.logger.events[-1][0][2] == "report_action"


def _action(web, report_id, action):
    return _hidden_of(web, f"/berichte/{report_id}", f"/berichte/{report_id}/aktion", action=action)


def _hidden_of(web, path, target, **match):
    html = web.browser.request("GET", path).body.decode()
    for target_path, block in re.findall(r'<form method="post" action="([^"]+)"[^>]*>(.*?)</form>', html, re.S):
        fields = {name: unescape(value) for name, value in
                  re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)">', block)}
        if target_path == target and all(fields.get(name) == value for name, value in match.items()):
            return fields
    raise AssertionError(f"no form {match} on {path}")


@pytest.mark.parametrize("status", list(ReportStatus))
def test_every_status_shows_its_actions(web, reports, status):
    report = filled(reports, status)
    status_code, html = page(web, f"/berichte/{report.id}")
    assert status_code == 200 and "Bericht löschen" in html or status in {ReportStatus.COUNTING,
                                                                          ReportStatus.RUNNING}


def test_report_page_lists_dates_unscheduled_and_rejected_mails(web, reports):
    report = filled(reports)
    status, html = page(web, f"/berichte/{report.id}")
    assert status == 200
    table = html[html.index("<tbody>"):html.index("</tbody>")]
    assert [title for title in ("Elternabend", "Basar", "Sitzung", "Fest", "Alt") if title in table] == [
        "Elternabend", "Basar", "Sitzung", "Fest"]
    assert "19:00–20:30" in table and "unsicher, außerhalb des Zeitraums" in table
    assert "von Hand ergänzt" in table and "korrigiert" in table
    assert "Relevante Mails ohne Termin" in html and "Als nicht relevant bewertet (1)" in html
    assert "Als xlsx-Liste herunterladen" in html and "(ohne Betreff)" not in table
    empty = reports.create(definition(title="Leer"), [], [])
    status, html = page(web, f"/berichte/{empty.id}")
    assert "Noch keine Termine erkannt." in html and "xlsx" not in html and "ohne Termin" not in html


def test_mail_pages_show_verdicts_and_escape_mail_content(web, reports):
    report = filled(reports)
    reports.store.save_mail(report.id, mail(2, "<script>alert(1)</script>", text="<img src=x onerror=alert(1)>",
                                            attachments=["protokoll.pdf"]))
    expected = {0: ["KI, Sicherheit: low", "Beleg: „am 20.01. um 19 Uhr“", "Wieder aufnehmen", "Nicht relevant"],
                1: ["Anhänge: protokoll.pdf", "Termin hinzufügen"],
                2: ["Nicht relevant (KI", "Doch relevant"],
                3: ["Übersprungen: ignorierter Absender."],
                4: ["Noch nicht bewertet.", "(ohne Betreff)"],
                5: ["Relevant (von dir festgelegt)", "Aufbewahrungsfrist entfernt"]}
    for index, texts in expected.items():
        status, html = page(web, f"/berichte/{report.id}/mail/{index}")
        assert status == 200 and all(text in html for text in texts), (index, html)
    raw = web.browser.request("GET", f"/berichte/{report.id}/mail/1").body.decode()
    assert "<script>alert" not in raw and "<img src=x" not in raw
    assert "Termin hinzufügen" not in page(web, f"/berichte/{report.id}/mail/3")[1]
    assert page(web, f"/berichte/{report.id}/mail/6")[0] == 404
    assert page(web, f"/berichte/{'b' * 32}/mail/0")[0] == 404


def test_corrections_change_relevance_and_dates(web, reports):
    report = filled(reports)
    base = f"/berichte/{report.id}/mail"
    # "Doch relevant" on the report page queues the date extraction.
    form = _hidden_of(web, f"/berichte/{report.id}", f"{base}/2", mode="relevant")
    assert post(web, f"{base}/2", form).headers["Location"] == f"{base}/2"
    assert reports.store.load(report.id).pending_extractions == ["INBOX:7:3"]
    post(web, f"{base}/2", _hidden_of(web, f"{base}/2", f"{base}/2", mode="irrelevant"))
    assert reports.store.mail(report.id, "INBOX:7:3").user_relevant is False
    post(web, f"{base}/0", _hidden_of(web, f"{base}/0", f"{base}/0", mode="exclude", event="e1"))
    post(web, f"{base}/0", _hidden_of(web, f"{base}/0", f"{base}/0", mode="include", event="e4"))
    edit = _hidden_of(web, f"{base}/0", f"{base}/0", mode="event", event="e2")
    response = post(web, f"{base}/0", edit | {"title": "Basar 2", "day": "2026-01-24", "start_time": "10:00",
                                              "end_time": "", "location": ""})
    assert response.status == 303
    add = _hidden_of(web, f"{base}/1", f"{base}/1", mode="event")
    post(web, f"{base}/1", add | {"title": "Nachtrag", "day": "2026-01-26", "start_time": "", "end_time": "",
                                  "location": "Schule"})
    events = {event.id: event for event in reports.store.mail(report.id, "INBOX:7:1").events}
    assert events["e1"].excluded and not events["e4"].excluded
    assert (events["e2"].title, events["e2"].day, events["e2"].start_time) == ("Basar 2", date(2026, 1, 24), "10:00")
    added = reports.store.mail(report.id, "INBOX:7:2").events
    assert [(event.title, event.manual, event.location) for event in added] == [("Nachtrag", True, "Schule")]


@pytest.mark.parametrize("fields, status", [
    ({"mode": "event", "title": "X", "day": "2026-02-30"}, 422),
    ({"mode": "event", "title": "X", "day": "2026-01-20", "start_time": "25:00"}, 422),
    ({"mode": "event", "title": " ", "day": "2026-01-20"}, 422),
    ({"mode": "gibt-es-nicht"}, 422),
    ({"mode": "exclude", "event": "e9"}, 409),
    ({"mode": "relevant", "version": "0"}, 409),
])
def test_invalid_corrections_change_nothing(web, reports, fields, status):
    report = filled(reports)
    form = _hidden_of(web, f"/berichte/{report.id}/mail/0", f"/berichte/{report.id}/mail/0", mode="irrelevant")
    assert post(web, f"/berichte/{report.id}/mail/0", form | fields).status == status
    assert reports.store.load(report.id).version == report.version


def test_skipped_mails_cannot_be_marked_relevant_and_missing_reports_answer_404(web, reports):
    report = filled(reports)
    fields = {"csrf": csrf(web), "key": "key-12345678", "version": str(report.version), "mode": "relevant"}
    assert post(web, f"/berichte/{report.id}/mail/3", fields).status == 409
    assert post(web, f"/berichte/{report.id}/mail/9", fields).status == 404
    assert post(web, f"/berichte/{'c' * 32}/mail/0", fields).status == 404
    assert page(web, f"/berichte/{'c' * 32}/export.xlsx")[0] == 404


def test_export_is_a_private_attachment(web, reports):
    report = filled(reports)
    web.site.clock = lambda: datetime(2026, 10, 11, 23, 30, tzinfo=UTC)
    response = web.browser.request("GET", f"/berichte/{report.id}/export.xlsx")
    headers = dict(response.all_headers())
    assert response.status == 200 and headers["Content-Type"].endswith("spreadsheetml.sheet")
    # 23:30 UTC is already the next day in Berlin.
    assert headers["Content-Disposition"] == 'attachment; filename="Taetigkeitsbericht-Elternbeirat-2026-10-12.xlsx"'
    assert headers["Cache-Control"] == "no-store" and headers["X-Content-Type-Options"] == "nosniff"
    book = load_workbook(io.BytesIO(response.body))
    assert book.sheetnames == ["Termine", "Ohne Termin", "Bericht"]
    web.site.activity.with_unscheduled = False
    response = web.browser.request("GET", f"/berichte/{report.id}/export.xlsx")
    assert load_workbook(io.BytesIO(response.body)).sheetnames == ["Termine", "Bericht"]
    assert web.logger.events[-1][0][2] == "report_exported"


def test_failed_model_answers_ask_the_user_to_decide(web, reports):
    report = reports.create(definition(), [], [])
    reports.store.save_mail(report.id, mail(1, "Ohne Urteil", relevant=None, confidence=None, reason="",
                                            events_extracted=False, failed_step="relevance"))
    reports.store.save_mail(report.id, mail(2, "Ohne Termine", failed_step="events"))
    status, html = page(web, f"/berichte/{report.id}/mail/0")
    assert "Die KI-Antwort war unbrauchbar. Bitte selbst entscheiden." in html and ">Doch relevant</button>" in html
    status, html = page(web, f"/berichte/{report.id}/mail/1")
    assert "Bitte von Hand ergänzen." in html and ">Nicht relevant</button>" in html
    form = _hidden_of(web, f"/berichte/{report.id}/mail/0", f"/berichte/{report.id}/mail/0", mode="relevant")
    post(web, f"/berichte/{report.id}/mail/0", form)
    assert "Relevant (von dir festgelegt)" in page(web, f"/berichte/{report.id}/mail/0")[1]


def test_excluded_view_marks_cancelled_and_excluded_dates(web, reports):
    report = reports.create(definition(), [], [])
    reports.store.save_mail(report.id, mail(1, "Absage", [
        ReportEvent(id="e1", title="Treffen", day=date(2026, 1, 20), cancelled=True),
        ReportEvent(id="e2", title="Ersatz", day=date(2026, 1, 21), excluded=True)]))
    status, html = page(web, f"/berichte/{report.id}")
    assert "Noch keine Termine erkannt." in html
    status, html = page(web, f"/berichte/{report.id}?excluded=on")
    table = html[html.index("<tbody>"):html.index("</tbody>")]
    assert "abgesagt" in table and "ausgeschlossen" in table


@pytest.mark.parametrize("value", ["../activity-report-x", "kurz", "A" * 32])
def test_progress_signal_rejects_malformed_report_ids(web, reports, value):
    # Regression: an id outside the report format reached the state store and failed with a server error.
    response = web.browser.request("GET", f"/api/stand?bericht={value}")
    assert response.status == 404


def test_a_report_reads_the_mailbox_chosen_for_it(web, reports):
    """Planung 3.1: with several mailboxes the form offers each one; folders are checked per mailbox."""
    web.site.activity.mailboxes = lambda: [(1, "Privat"), (2, "Arbeit")]
    web.site.activity.suggest_folders = lambda mailbox: {1: ["INBOX"], 2: ["Projekte"]}[mailbox]
    web.site.activity.available_folders = lambda mailbox: {1: ["INBOX"], 2: ["INBOX", "Projekte"]}[mailbox]
    status, html = page(web, "/berichte/neu?postfach=2")
    assert status == 200 and '<a href="/berichte/neu?postfach=1">Privat</a>' in html
    assert "<strong>Arbeit</strong>" in html and "Projekte</textarea>" in html
    assert 'name="mailbox" value="2"' in html
    for query in ("?postfach=7", "?postfach=x", ""):
        assert 'name="mailbox" value="1"' in page(web, "/berichte/neu" + query)[1]
    fields = {**new_form(web), **NEW, "folders": "Projekte", "mailbox": "2"}
    assert post(web, "/berichte", fields).status == 303
    (report,) = reports.store.reports()
    assert report.definition.mailbox == 2
    # Folders of another mailbox are refused, the form keeps the chosen mailbox.
    refused = post(web, "/berichte", {**new_form(web), **NEW, "folders": "Projekte", "mailbox": "1"})
    assert refused.status == 422 and 'name="mailbox" value="1"' in refused.body.decode()


def test_a_single_mailbox_shows_no_mailbox_choice(web, reports):
    status, html = page(web, "/berichte/neu")
    assert status == 200 and "Postfach:" not in html and 'name="mailbox" value="1"' in html
