"""xlsx export of an activity report: read back, dates, no formulas, limits."""
from __future__ import annotations

import io
from datetime import UTC, date, datetime

import pytest
from activity_fakes import definition
from openpyxl import load_workbook

from mailhelp import __version__
from mailhelp.activity.export import CELL_LIMIT, REMOVED, filename, safe_text, workbook
from mailhelp.activity.models import ActivityReport, ReportEvent, ReportMail

RECEIVED = datetime(2026, 1, 10, 8, 30, tzinfo=UTC)
ATTACKS = ["=HYPERLINK(\"https://evil.test\",\"klick\")", "+cmd|' /C calc'!A0", "-2+3", "@SUM(A1:A2)",
           "\tTab", "\rCR"]


def event(identifier, day, title="Elternabend", **changes):
    return ReportEvent(id=identifier, title=title, day=day, **changes)


def mail(key, subject, events=(), **changes):
    values = dict(key=key, folder="INBOX", uid=int(key), received_at=RECEIVED, sender="Schule <info@schule.test>",
                  subject=subject, text="Volltext äöü", relevant=True, events=list(events))
    return ReportMail(**(values | changes))


def report():
    return ActivityReport(id="a" * 32, definition=definition(title="Elternbeirat 4b / 2026"),
                          folders=[{"folder": "INBOX", "uids": [1]}])


def read(data):
    return load_workbook(io.BytesIO(data))


def test_dates_numbers_and_sheets():
    mails = [
        mail("1", "Einladung", [event("e1", date(2026, 1, 20), start_time="19:00", end_time="20:30", location="Aula"),
                                event("e2", date(2026, 1, 27), cancelled=True),
                                event("e3", date(2026, 1, 5), title="Vorgespräch", excluded=True)]),
        mail("2", "Erinnerung", [event("e1", date(2026, 1, 20), title="elternabend ", start_time="19:00")],
             text_removed=True, text=""),
        mail("3", "Protokoll"),
        mail("4", "Werbung", [event("e1", date(2026, 1, 3))], relevant=False),
        mail("5", "Übersprungen", skipped="sender", relevant=None),
        mail("6", "Frühtermin", [event("e1", date(2026, 1, 15), title="Frühtermin", start_time="08:00", manual=True)],
             relevant=False, user_relevant=True),
    ]
    book = read(workbook(report(), mails, "Europe/Berlin", True))
    assert book.sheetnames == ["Termine", "Ohne Termin", "Bericht"]
    rows = list(book["Termine"].iter_rows(values_only=True))
    assert rows[0] == ("Termin-Nr.", "Termin", "Datum", "Uhrzeit", "Ort", "Mail-Betreff", "Absender",
                       "Mail-Datum", "Volltext")
    assert [(row[0], row[1], row[3], row[5], row[8]) for row in rows[1:]] == [
        (1, "Frühtermin", "08:00", "Frühtermin", "Volltext äöü"),
        (2, "Elternabend", "19:00–20:30", "Einladung", "Volltext äöü"),
        (2, "elternabend ", "19:00", "Erinnerung", REMOVED)]
    assert rows[1][2] == datetime(2026, 1, 15) and rows[1][7] == datetime(2026, 1, 10, 9, 30)
    assert book["Termine"].cell(2, 3).number_format == "DD.MM.YYYY" and book["Termine"].freeze_panes == "A2"
    assert list(book["Ohne Termin"].iter_rows(values_only=True))[1][0] == "Protokoll"
    summary = dict(book["Bericht"].iter_rows(min_row=2, values_only=True))
    assert summary["Zeitraum"] == "01.01.2026 – 31.01.2026" and summary["Termine"] == "3"
    assert summary["Mailhelp-Version"] == __version__ and summary["Hinweis"].endswith("geprüft: ja")
    assert summary["Ordner"] == "INBOX"
    plain = read(workbook(report().model_copy(update={"folders": []}), [mail("3", "Protokoll")], "UTC", False))
    assert plain.sheetnames == ["Termine", "Bericht"] and plain["Termine"].max_row == 1
    plain_summary = dict(plain["Bericht"].iter_rows(min_row=2, values_only=True))
    assert plain_summary["Hinweis"].endswith("geprüft: nein") and plain_summary["Ordner"] == "INBOX"


@pytest.mark.parametrize("attack", ATTACKS)
def test_mail_content_never_becomes_a_formula(attack):
    # Dates refuse control characters already, so CR reaches only the mail fields.
    field = attack.replace("\r", "")
    mails = [mail("1", attack, [event("e1", date(2026, 1, 20), title=field, location=field)],
                  sender=attack, text=attack)]
    book = read(workbook(report(), mails, "UTC", True))
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@", "\t", "\r"))
    # A leading CR becomes a line break first, which no spreadsheet reads as a formula.
    expected_prefix = "\n" if attack.startswith("\r") else "'"
    assert book["Termine"].cell(2, 6).value.startswith(expected_prefix)


def test_cells_are_cleaned_and_bounded():
    assert safe_text("a\x00b\x07c\r\nd\re\tf") == "abc\nd\ne\tf"
    assert len(safe_text("x" * 40000)) == CELL_LIMIT
    assert filename(report(), date(2026, 10, 12)) == "Taetigkeitsbericht-Elternbeirat-4b-2026-2026-10-12.xlsx"
    unnamed = report().model_copy(update={"definition": definition(title="!!!")})
    assert filename(unnamed, date(2026, 10, 12)) == "Taetigkeitsbericht-Bericht-2026-10-12.xlsx"
